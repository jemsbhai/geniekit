"""Compile the six public GENIE BPC Variable Synopsis spreadsheets into one table.

AACR publishes a "Variable Synopsis" workbook for every public BPC release
(linked from the cohort pages on aacrprojectgenie.org). Each workbook lists,
per analytic dataset, the variable name, field label, data type and permitted
values. This script downloads the six workbooks into ``data-raw/cache`` (kept
out of git), normalises them into one long table and writes
``src/geniekit/resources/dictionary.csv.gz``.

Run from the repository root::

    python data-raw/build_dictionary.py

The output is metadata only (variable names, labels, types, value lists). It
contains no patient-level data.
"""

from __future__ import annotations

import csv
import datetime as dt
import gzip
import json
import re
import sys
import urllib.request
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - build-time dependency only
    sys.exit("openpyxl is required to build the dictionary: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
RELEASES = ROOT / "src" / "geniekit" / "resources" / "releases.json"
CACHE = ROOT / "data-raw" / "cache"
OUT = ROOT / "src" / "geniekit" / "resources" / "dictionary.csv.gz"

# Synopsis "Dataset" column -> geniekit logical table name. The synopsis lists
# one "Cancer diagnosis dataset"; it describes both the index and the
# non-index cancer files, which share a schema.
DATASET_TO_TABLE = {
    "Patient-level dataset": "patient",
    "Cancer diagnosis dataset": "cancer",
    "Cancer-directed regimen dataset": "regimen",
    "PRISSMM Imaging level dataset": "imaging",
    "PRISSMM Pathology level dataset": "pathology",
    "PRISSMM Medical Oncologist Assessment level dataset": "med_onc",
    "PRISSMM Tumor Marker level dataset": "tumor_marker",
    "Cancer-Directed Radiation Therapy dataset": "radiation",
    "Cancer panel test level dataset": "cpt",
}

COLUMNS = [
    "cohort",
    "version",
    "schema_generation",
    "table",
    "dataset_label",
    "variable",
    "label",
    "dtype",
    "values",
    "n_values",
    "order",
]

# bullet (U+2022), middle dot (U+00B7) and hyphen used as list markers in the Values cells
_BULLET_CHARS = "".join(chr(c) for c in (0x2022, 0x00B7)) + "-"
_BULLET_RE = re.compile("^[" + re.escape(_BULLET_CHARS) + r"]+\s*")


def _clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("_x000D_", "").replace("\r", "")
    return text.strip()


def _parse_values(raw: str) -> list[str]:
    """Split the bullet list in the synopsis "Values" cell into clean strings."""
    if not raw:
        return []
    out: list[str] = []
    for line in raw.split("\n"):
        line = line.strip()
        line = _BULLET_RE.sub("", line).strip()
        line = line.strip("\t ").strip()
        if line:
            out.append(line)
    return out


def _download(url: str, dest: Path) -> Path:
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    dest.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(url, headers={"User-Agent": "geniekit-build/0.1"})
    with urllib.request.urlopen(req, timeout=120) as resp, open(dest, "wb") as fh:
        fh.write(resp.read())
    return dest


def parse_workbook(path: Path, cohort: str, version: str, generation: int) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = wb["Variable Synopsis"] if "Variable Synopsis" in wb.sheetnames else wb.worksheets[-1]
    rows = [tuple(_clean(c) for c in r) for r in sheet.iter_rows(values_only=True)]
    header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Dataset" and "Variable Name" in r)
    current_dataset = ""
    order = 0
    records: list[dict] = []
    for r in rows[header_idx + 1 :]:
        r = list(r) + [""] * 5
        dataset, variable, label, dtype, values = r[:5]
        if dataset:
            current_dataset = dataset
        if not variable:
            continue
        if current_dataset not in DATASET_TO_TABLE:
            raise ValueError(f"{cohort}: unknown dataset label {current_dataset!r}")
        order += 1
        levels = _parse_values(values)
        records.append(
            {
                "cohort": cohort,
                "version": version,
                "schema_generation": generation,
                "table": DATASET_TO_TABLE[current_dataset],
                "dataset_label": current_dataset,
                "variable": variable,
                "label": label,
                "dtype": dtype,
                "values": json.dumps(levels, ensure_ascii=False),
                "n_values": len(levels),
                "order": order,
            }
        )
    return records


def main() -> None:
    releases = json.loads(RELEASES.read_text(encoding="utf-8"))["releases"]
    all_records: list[dict] = []
    for rel in releases:
        url = rel["sources"]["variable_synopsis_url"]
        dest = CACHE / f"{rel['cohort']}_{rel['version']}_variable_synopsis.xlsx"
        _download(url, dest)
        recs = parse_workbook(dest, rel["cohort"], rel["version"], rel["schema_generation"])
        print(f"{rel['label']:22s} {len(recs):5d} variables from {dest.name}")
        all_records.extend(recs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(OUT, "wt", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(all_records)
    meta = {
        "generated_on": dt.date.today().isoformat(),
        "generator": "data-raw/build_dictionary.py",
        "n_rows": len(all_records),
        "sources": {rel["label"]: rel["sources"]["variable_synopsis_url"] for rel in releases},
    }
    (OUT.parent / "dictionary.meta.json").write_text(json.dumps(meta, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {OUT} ({len(all_records)} rows)")


if __name__ == "__main__":
    main()
